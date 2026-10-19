import os
from io import BytesIO
from datetime import datetime
from flask import Flask, render_template, request, send_file, flash
from dotenv import load_dotenv
from generator import generate_material

load_dotenv()
app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "change-this-secret-key")

MATERIAL_TYPES = {
    "soal": "Paket Soal",
    "kisi_kisi": "Kisi-kisi Penilaian",
    "lkpd": "Lembar Kerja Peserta Didik",
    "rubrik": "Rubrik Penilaian",
    "bahan_ajar": "Bahan Ajar",
}
SUBJECTS = [
    "Matematika", "Bahasa Indonesia", "IPAS", "Pendidikan Pancasila",
    "Bahasa Inggris", "Informatika", "IPA", "IPS", "PJOK", "Seni Budaya",
]
GRADES = ["Kelas 1", "Kelas 2", "Kelas 3", "Kelas 4", "Kelas 5", "Kelas 6",
          "Kelas 7", "Kelas 8", "Kelas 9", "Kelas 10", "Kelas 11", "Kelas 12"]

@app.get("/")
def index():
    return render_template("index.html", material_types=MATERIAL_TYPES,
                           subjects=SUBJECTS, grades=GRADES)

@app.post("/generate")
def generate():
    material_type = request.form.get("material_type", "soal")
    subject = request.form.get("subject", "Matematika").strip()
    grade = request.form.get("grade", "Kelas 4")
    topic = request.form.get("topic", "").strip()
    curriculum = request.form.get("curriculum", "Kurikulum Merdeka")
    difficulty = request.form.get("difficulty", "Sedang")
    count_raw = request.form.get("count", "10")
    learning_objective = request.form.get("learning_objective", "").strip()
    teacher_notes = request.form.get("teacher_notes", "").strip()

    if material_type not in MATERIAL_TYPES:
        flash("Jenis dokumen tidak valid.", "error")
        return render_template("index.html", material_types=MATERIAL_TYPES,
                               subjects=SUBJECTS, grades=GRADES), 400
    if not topic:
        flash("Isi topik atau materi terlebih dahulu.", "error")
        return render_template("index.html", material_types=MATERIAL_TYPES,
                               subjects=SUBJECTS, grades=GRADES), 400
    try:
        count = max(1, min(int(count_raw), 30))
    except ValueError:
        count = 10

    payload = {
        "material_type": material_type,
        "material_label": MATERIAL_TYPES[material_type],
        "subject": subject,
        "grade": grade,
        "topic": topic,
        "curriculum": curriculum,
        "difficulty": difficulty,
        "count": count,
        "learning_objective": learning_objective,
        "teacher_notes": teacher_notes,
    }
    result = generate_material(payload)
    return render_template("result.html", payload=payload, result=result,
                           generated_at=datetime.now().strftime("%d-%m-%Y %H:%M"))

@app.post("/download")
def download():
    title = request.form.get("title", "Dokumen GuruCerdas AI")
    content = request.form.get("content", "")
    if not content.strip():
        flash("Tidak ada konten untuk diunduh.", "error")
        return render_template("index.html", material_types=MATERIAL_TYPES,
                               subjects=SUBJECTS, grades=GRADES), 400
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(31, 78, 121)
        doc.add_paragraph("Dibuat dengan GuruCerdas AI • Periksa dan sesuaikan sebelum digunakan.")
        for line in content.splitlines():
            line = line.rstrip()
            if not line:
                doc.add_paragraph()
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif len(line) > 2 and line[0].isdigit() and line[1:3] in (". ", ") "):
                doc.add_paragraph(line[3:])
            else:
                doc.add_paragraph(line)
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        filename = "".join(c for c in title if c.isalnum() or c in (" ", "-", "_")).strip().replace(" ", "_")
        return send_file(output, as_attachment=True,
                         download_name=f"{filename[:70] or 'gurucerdas'}.docx",
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except ImportError:
        from io import BytesIO
        output = BytesIO(content.encode("utf-8"))
        return send_file(output, as_attachment=True, download_name="gurucerdas.txt",
                         mimetype="text/plain; charset=utf-8")

@app.get("/health")
def health():
    return {"status": "ok", "app": "GuruCerdas AI"}

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=os.getenv("FLASK_DEBUG", "false").lower() == "true")
